from decimal import Decimal
from io import BytesIO
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.conf import settings
from django.core.mail import EmailMultiAlternatives
from django.contrib.staticfiles import finders
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.utils import timezone
from django.db.models import F, Q, Count, Sum
from django.db import transaction
from django.views.decorators.http import require_POST

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image

from .models import (
    InventoryReception, InventoryReceptionLine, WorkshopDelivery,
    InventoryExit, InventoryExitLine, DispatchRemission, DispatchRemissionLine, RemissionSequence,
)
from auditoria.utils import registrar_movimiento


def inventario_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        permitido = (
            request.user.is_authenticated
            and (
                request.user.is_superuser
                or request.user.groups.filter(name__iexact="INVENTARIO").exists()
            )
        )

        if not permitido:
            messages.error(
                request,
                "No tienes permiso para acceder al módulo de Inventario."
            )
            return redirect("/")

        return view_func(request, *args, **kwargs)

    return wrapper


@login_required
@inventario_required
def inventario_dashboard(request):
    # ======================================================
    # RECEPCIONES: solo pendientes o parciales
    # ======================================================
    recepciones_qs = (
        InventoryReception.objects
        .select_related("purchase_request", "creado_por")
        .prefetch_related("lineas")
        .order_by("-actualizado_en")
    )

    recepciones_pendientes = []
    recepciones_completas = []

    for r in recepciones_qs:
        lineas = list(r.lineas.all())
        total = len(lineas)
        listas = 0
        parciales = 0

        for linea in lineas:
            esperada = Decimal(linea.cantidad_esperada or 0)
            recibida = Decimal(linea.cantidad_recibida or 0)

            if esperada > 0 and recibida >= esperada:
                listas += 1
            elif recibida > 0:
                parciales += 1

        r.total_lineas = total
        r.lineas_listas = listas
        r.lineas_parciales = parciales
        r.recepcion_completa = total > 0 and listas == total

        if total > 0:
            r.porcentaje = round((listas / total) * 100)
        else:
            r.porcentaje = 0

        if r.recepcion_completa:
            recepciones_completas.append(r)
        else:
            recepciones_pendientes.append(r)

    # ======================================================
    # ENTREGAS: separar taller activo vs historial
    # ======================================================
    entregas_qs = (
        WorkshopDelivery.objects
        .select_related("purchase_request", "creado_por")
        .prefetch_related("lineas")
        .order_by("-actualizado_en")
    )

    entregas_taller_pendientes = []
    historial_entregas = []

    for entrega in entregas_qs:
        lineas = list(entrega.lineas.all())
        total_lineas = len(lineas)
        lineas_completas = 0
        cantidad_requerida_total = Decimal("0")
        cantidad_entregada_total = Decimal("0")

        for linea in lineas:
            requerida = Decimal(linea.cantidad_requerida or 0)
            entregada = Decimal(linea.cantidad_entregada or 0)

            cantidad_requerida_total += requerida
            if requerida > 0:
                cantidad_entregada_total += min(entregada, requerida)

            if requerida <= 0 or entregada >= requerida:
                lineas_completas += 1

        entrega.total_lineas = total_lineas
        entrega.lineas_completas = lineas_completas
        entrega.entrega_completa = (
            total_lineas > 0
            and lineas_completas == total_lineas
        )

        if cantidad_requerida_total > 0:
            entrega.porcentaje_entrega = min(
                100,
                round(float(
                    (cantidad_entregada_total / cantidad_requerida_total) * 100
                ))
            )
        else:
            entrega.porcentaje_entrega = 0

        destino = str(getattr(entrega, "destino", "TALLER") or "TALLER").upper()
        entrega.destino_codigo = destino

        # TALLER:
        # - pendiente/parcial permanece en el bloque operativo.
        # - completa pasa a historial.
        #
        # CAMPO / INVENTARIO (cliente):
        # - se consideran salida de Inventario y se muestran en historial,
        #   porque Inventario ya no tiene una entrega a Taller que gestionar.
        if destino == "TALLER" and not entrega.entrega_completa:
            entregas_taller_pendientes.append(entrega)
        else:
            historial_entregas.append(entrega)

    # ======================================================
    # BOM / SOLICITUDES PENDIENTES DE REVISIÓN DE INVENTARIO
    # La solicitud puede existir técnicamente, pero Compras no la ve hasta
    # que Inventario confirme cantidades disponibles.
    # ======================================================
    from compras_oil.models import PurchaseRequest

    revisiones_pendientes = (
        PurchaseRequest.objects
        .filter(inventario_revisado_en__isnull=True)
        .exclude(estado="CERRADA")
        .exclude(bom__workorder__paw__estado_operativo__in=["FACTURADO", "RADICADO"])
        .select_related("bom", "bom__workorder", "creado_por")
        .annotate(
            total_lineas_bom=Count(
                "lineas",
                filter=Q(lineas__cantidad_requerida__gt=0),
                distinct=True,
            )
        )
        .order_by("creado_en")
    )

    # PAW revisados que ya tienen todo el material necesario disponible
    # (por stock o por recepción de compra) y aún no tienen entrega generada.
    candidatos_entrega = (
        PurchaseRequest.objects
        .filter(inventario_revisado_en__isnull=False)
        .exclude(estado="CERRADA")
        .exclude(bom__workorder__paw__estado_operativo__in=["FACTURADO", "RADICADO"])
        .select_related("bom", "bom__workorder")
        .prefetch_related("lineas", "recepcion_inventario__lineas")
        .order_by("actualizado_en")
    )
    listos_para_entrega = []
    for compra in candidatos_entrega:
        try:
            compra.entrega_taller
            continue
        except Exception:
            pass
        if _material_comprado_completamente_recibido(compra):
            listos_para_entrega.append(compra)

    return render(request, "inventario/dashboard.html", {
        "revisiones_pendientes": revisiones_pendientes,
        "total_revisiones_pendientes": revisiones_pendientes.count(),
        "listos_para_entrega": listos_para_entrega,
        "total_listos_para_entrega": len(listos_para_entrega),
        "recepciones": recepciones_pendientes,
        "recepciones_completas": recepciones_completas,
        "entregas": entregas_taller_pendientes,
        "historial_entregas": historial_entregas,

        "total_recepciones": len(recepciones_pendientes),
        "total_recepciones_completas": len(recepciones_completas),
        "total_entregas": len(entregas_taller_pendientes),
        "total_historial_entregas": len(historial_entregas),

        "lineas_pendientes": InventoryReceptionLine.objects.filter(
            estado="PENDIENTE"
        ).count(),
        "lineas_parciales": InventoryReceptionLine.objects.filter(
            estado="PARCIAL"
        ).count(),
        "lineas_listas": InventoryReceptionLine.objects.filter(
            estado="LISTO"
        ).count(),
    })

@login_required
@inventario_required
def revision_bom_detail(request, pk):
    """Revisa únicamente el delta pendiente del BOM y crea reservas sin mover existencia física."""
    from compras_oil.models import PurchaseRequest
    from .models import InventoryStock, InventoryReservation

    compra = get_object_or_404(
        PurchaseRequest.objects.select_related("bom", "bom__workorder", "creado_por")
        .prefetch_related("lineas__bom_item"), pk=pk,
    )

    # CLAVE DEL FLUJO INCREMENTAL:
    # una línea aparece solo si el BOM requiere más de lo que Inventario ya revisó.
    todas_lineas = compra.lineas.filter(cantidad_requerida__gt=0).order_by("id")
    lineas = []
    for linea in todas_lineas:
        requerida_total = Decimal(linea.cantidad_requerida or 0)
        revisada = Decimal(linea.cantidad_revisada_inventario or 0)
        pendiente = max(requerida_total - revisada, Decimal("0"))
        if pendiente > 0:
            linea.cantidad_pendiente_inventario = pendiente
            lineas.append(linea)

    empresa_stock = InventoryStock.Empresa.IMPETUS

    if request.method == "POST":
        if compra.inventario_revisado_en:
            messages.info(request, "Este BOM ya fue revisado por Inventario.")
            return redirect("inventario:revision_bom_detail", pk=compra.pk)

        if not lineas:
            messages.info(request, "No hay requerimientos nuevos o incrementales pendientes de revisión.")
            return redirect("inventario:revision_bom_detail", pk=compra.pk)

        errores, cantidades = [], {}
        with transaction.atomic():
            for linea in lineas:
                pendiente = Decimal(linea.cantidad_pendiente_inventario or 0)
                raw = (request.POST.get(f"cantidad_disponible_{linea.id}") or "0").strip()
                try:
                    solicitada = Decimal(raw.replace(",", "."))
                except Exception:
                    errores.append(f"Cantidad inválida para {linea.codigo or linea.descripcion}.")
                    continue
                if solicitada < 0 or solicitada > pendiente:
                    errores.append(
                        f"La reserva de {linea.codigo or linea.descripcion} debe estar entre 0 y {pendiente}."
                    )
                    continue

                stock = InventoryStock.objects.select_for_update().filter(
                    empresa=empresa_stock, codigo__iexact=(linea.codigo or "").strip()
                ).first()
                disponible_real = Decimal(stock.cantidad_disponible) if stock else Decimal("0")
                if solicitada > disponible_real:
                    errores.append(
                        f"{linea.codigo or linea.descripcion}: intentas reservar {solicitada}, "
                        f"pero el disponible real es {disponible_real}."
                    )
                    continue
                cantidades[linea.id] = (solicitada, stock)

            if errores:
                transaction.set_rollback(True)
            else:
                for linea in lineas:
                    cantidad, stock = cantidades.get(linea.id, (Decimal("0"), None))
                    pendiente = Decimal(linea.cantidad_pendiente_inventario or 0)

                    # cantidad_disponible es acumulada para esta PurchaseLine.
                    linea.cantidad_disponible = Decimal(linea.cantidad_disponible or 0) + cantidad
                    linea.cantidad_revisada_inventario = (
                        Decimal(linea.cantidad_revisada_inventario or 0) + pendiente
                    )
                    linea.save(update_fields=[
                        "cantidad_disponible",
                        "cantidad_revisada_inventario",
                        "cantidad_a_comprar",
                    ])

                    if cantidad > 0 and stock:
                        InventoryReservation.objects.create(
                            stock=stock,
                            cantidad=cantidad,
                            purchase_request=compra,
                            purchase_line=linea,
                            creado_por=request.user,
                            observacion=f"Reserva BOM PAW #{compra.paw_numero}",
                        )
                        stock.cantidad_reservada = F("cantidad_reservada") + cantidad
                        stock.save(update_fields=["cantidad_reservada", "actualizado_en"])

                compra.inventario_revisado_en = timezone.now()
                compra.inventario_revisado_por = request.user
                compra.save(update_fields=["inventario_revisado_en", "inventario_revisado_por", "actualizado_en"])

                registrar_movimiento(
                    request=request, paw_numero=compra.paw_numero, modulo="INVENTARIO",
                    accion="BOM revisado y reservado",
                    descripcion="Inventario confirmó únicamente los requerimientos nuevos/incrementales del PAW.",
                    objeto=compra,
                )
                faltantes = compra.lineas.filter(cantidad_a_comprar__gt=0).exists()
                try:
                    paw = compra.bom.workorder.paw
                    paw.estado_operativo = "EN_COMPRAS" if faltantes else "MATERIAL_RECIBIDO"
                    paw.save(update_fields=["estado_operativo"])
                except Exception:
                    pass

        if errores:
            for error in errores:
                messages.error(request, error)
        else:
            if faltantes:
                messages.success(request, "Requerimientos pendientes revisados. Solo los faltantes quedan para Compras.")
                return redirect("inventario:dashboard")
            messages.success(request, "Los nuevos requerimientos quedaron reservados. No se requiere compra adicional.")
            return redirect("inventario:revision_bom_detail", pk=compra.pk)

    # GET: el stock se consulta después de determinar qué delta está pendiente.
    for linea in lineas:
        stock = InventoryStock.objects.filter(
            empresa=empresa_stock, codigo__iexact=(linea.codigo or "").strip()
        ).first()
        linea.stock_fisico = Decimal(stock.cantidad_fisica) if stock else Decimal("0")
        linea.stock_reservado = Decimal(stock.cantidad_reservada) if stock else Decimal("0")
        linea.stock_disponible_real = Decimal(stock.cantidad_disponible) if stock else Decimal("0")
        pendiente = Decimal(linea.cantidad_pendiente_inventario or 0)
        linea.reserva_sugerida = min(pendiente, linea.stock_disponible_real)
        linea.reserva_sugerida_input = format(linea.reserva_sugerida, "f")
        linea.faltante_sugerido = max(pendiente - linea.reserva_sugerida, Decimal("0"))

    total_requerido = sum((Decimal(x.cantidad_pendiente_inventario or 0) for x in lineas), Decimal("0"))
    total_disponible = sum((x.reserva_sugerida for x in lineas), Decimal("0"))
    total_comprar = sum((x.faltante_sugerido for x in lineas), Decimal("0"))

    return render(request, "inventario/revision_bom_detail.html", {
        "compra": compra,
        "lineas": lineas,
        "total_requerido": total_requerido,
        "total_disponible": total_disponible,
        "total_comprar": total_comprar,
        "modo_incremental": any(Decimal(x.cantidad_revisada_inventario or 0) > 0 for x in todas_lineas),
    })


def _material_comprado_completamente_recibido(compra):
    """True si cada faltante que debía comprarse ya fue recibido por Inventario."""
    lineas_compra = list(compra.lineas.filter(cantidad_a_comprar__gt=0))
    if not lineas_compra:
        return True
    try:
        recepcion = compra.recepcion_inventario
    except Exception:
        return False
    recibidas = {x.purchase_line_id: x for x in recepcion.lineas.all()}
    for linea in lineas_compra:
        r = recibidas.get(linea.id)
        if not r:
            return False
        if Decimal(r.cantidad_recibida or 0) < Decimal(linea.cantidad_a_comprar or 0):
            return False
    return True


@require_POST
@login_required
@inventario_required
def generar_entrega(request, pk):
    """Inventario define el destino y genera la salida física del material del PAW."""
    from compras_oil.models import PurchaseRequest
    from .models import WorkshopDeliveryLine

    compra = get_object_or_404(
        PurchaseRequest.objects.prefetch_related("lineas", "recepcion_inventario__lineas"),
        pk=pk,
    )

    if not compra.inventario_revisado_en:
        messages.error(request, "Primero debes revisar el BOM en Inventario.")
        return redirect("inventario:revision_bom_detail", pk=compra.pk)

    if not _material_comprado_completamente_recibido(compra):
        messages.error(request, "Aún existen materiales comprados pendientes de recepción.")
        return redirect("inventario:dashboard")

    destino = (request.POST.get("destino") or "").upper().strip()
    destinos_validos = {"TALLER", "CAMPO", "CLIENTE", "INVENTARIO"}
    if destino not in destinos_validos:
        messages.error(request, "Selecciona Taller, Campo, Cliente/Despacho o Inventario/Bodega.")
        return redirect("inventario:dashboard")

    try:
        paw = compra.bom.workorder.paw
    except Exception:
        paw = None

    if destino == "TALLER" and paw and not getattr(paw, "aplica_taller", False):
        messages.error(request, "Este PAW no tiene habilitado Taller.")
        return redirect("inventario:dashboard")
    if destino == "CAMPO" and paw and not getattr(paw, "aplica_campo", False):
        messages.error(request, "Este PAW no tiene habilitado Campo.")
        return redirect("inventario:dashboard")
    if destino in {"CLIENTE", "INVENTARIO"} and paw:
        solo_compras = (
            getattr(paw, "aplica_compras", False)
            and not getattr(paw, "aplica_taller", False)
            and not getattr(paw, "aplica_campo", False)
        )
        if not solo_compras:
            messages.error(
                request,
                "Cliente/Despacho e Inventario/Bodega solo están disponibles cuando el PAW es únicamente de Compras/Materiales.",
            )
            return redirect("inventario:dashboard")

    with transaction.atomic():
        entrega, created = WorkshopDelivery.objects.get_or_create(
            purchase_request=compra,
            defaults={"creado_por": request.user, "destino": destino},
        )
        if not created and entrega.destino != destino:
            if entrega.lineas.filter(cantidad_entregada__gt=0).exists():
                messages.error(request, "No puedes cambiar el destino porque la entrega ya inició.")
                return redirect("inventario:entrega_taller_detail", pk=entrega.pk)
            entrega.destino = destino
            entrega.save(update_fields=["destino", "actualizado_en"])

        creadas = 0
        for linea in compra.lineas.filter(cantidad_requerida__gt=0):
            linea_entrega, nueva = WorkshopDeliveryLine.objects.get_or_create(
                delivery=entrega,
                purchase_line=linea,
                defaults={
                    "codigo": linea.codigo or "",
                    "descripcion": linea.descripcion or "",
                    "unidad": linea.unidad or "",
                    # Inventario entrega el total requerido por el BOM:
                    # stock disponible + material comprado.
                    "cantidad_requerida": Decimal(linea.cantidad_requerida or 0),
                },
            )
            # Si Taller agregó un requerimiento incremental después de crear la entrega,
            # actualizamos el total requerido sin tocar lo ya entregado.
            requerida_actual = Decimal(linea.cantidad_requerida or 0)
            if not nueva and Decimal(linea_entrega.cantidad_requerida or 0) != requerida_actual:
                linea_entrega.cantidad_requerida = requerida_actual
                linea_entrega.codigo = linea.codigo or linea_entrega.codigo
                linea_entrega.descripcion = linea.descripcion or linea_entrega.descripcion
                linea_entrega.unidad = linea.unidad or linea_entrega.unidad
                linea_entrega.save(update_fields=["cantidad_requerida", "codigo", "descripcion", "unidad"])
            creadas += int(nueva)

    messages.success(
        request,
        f"Entrega a {entrega.get_destino_display()} generada por Inventario. Líneas nuevas: {creadas}."
    )
    return redirect("inventario:entrega_taller_detail", pk=entrega.pk)


def _enviar_alerta_recepcion(recepcion, porcentaje, pendientes, umbral):
    """
    Envía una alerta inmediata cuando una recepción alcanza por primera vez
    el 80% o el 100%. Los destinatarios son COMPRAS, ALERTAS_TALLER y
    superusuarios activos con correo.
    """
    User = get_user_model()

    destinatarios = list(
        User.objects.filter(is_active=True)
        .filter(
            Q(groups__name="COMPRAS")
            | Q(groups__name="ALERTAS_TALLER")
            | Q(is_superuser=True)
        )
        .exclude(email="")
        .values_list("email", flat=True)
        .distinct()
    )

    if not destinatarios:
        return 0

    pr = recepcion.purchase_request
    paw_numero = getattr(pr, "paw_numero", None) or "-"
    paw_nombre = getattr(pr, "paw_nombre", "") or ""

    base_url = getattr(
        settings,
        "IMPETUS_CONTROL_URL",
        "https://www.impetuscontrol.com",
    ).rstrip("/")

    inventario_url = f"{base_url}/inventario/recepcion/{recepcion.pk}/"
    taller_url = f"{base_url}/taller/"

    if umbral == 100:
        asunto = f"IMPETUS CONTROL · PAW #{paw_numero} · Recepción 100% completa"
        titulo = "Recepción de material 100% completa"
        mensaje_estado = (
            "Inventario confirmó la recepción completa de los materiales "
            "asociados al PAW. Taller puede continuar con la programación "
            "correspondiente."
        )
    else:
        asunto = f"IMPETUS CONTROL · PAW #{paw_numero} · Recepción {porcentaje}%"
        titulo = "Recepción de material alcanzó el 80%"
        mensaje_estado = (
            "La recepción alcanzó al menos el 80%. Se recomienda a Taller "
            "evaluar si los componentes recibidos permiten iniciar o avanzar "
            "el ensamble."
        )

    pendientes_texto = ""
    pendientes_html = ""

    if pendientes:
        pendientes_texto = "\n\nMATERIALES AÚN PENDIENTES:\n"
        filas = []

        for item in pendientes:
            pendientes_texto += (
                f"- {item['codigo']} | {item['descripcion']} | "
                f"Pendiente: {item['faltante']} {item['unidad']}\n"
            )

            filas.append(
                "<tr>"
                f"<td style='padding:8px;border-bottom:1px solid #e5e7eb;'>{item['codigo']}</td>"
                f"<td style='padding:8px;border-bottom:1px solid #e5e7eb;'>{item['descripcion']}</td>"
                f"<td style='padding:8px;border-bottom:1px solid #e5e7eb;'>{item['faltante']}</td>"
                f"<td style='padding:8px;border-bottom:1px solid #e5e7eb;'>{item['unidad']}</td>"
                "</tr>"
            )

        pendientes_html = f"""
        <div style="margin-top:24px;">
            <div style="font-size:16px;font-weight:800;color:#0f172a;margin-bottom:10px;">
                Materiales aún pendientes
            </div>
            <table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"
                   style="width:100%;border-collapse:collapse;font-size:13px;border:1px solid #e2e8f0;">
                <thead>
                    <tr style="background:#f8fafc;">
                        <th align="left" style="padding:9px;border-bottom:1px solid #e2e8f0;">Código</th>
                        <th align="left" style="padding:9px;border-bottom:1px solid #e2e8f0;">Descripción</th>
                        <th align="left" style="padding:9px;border-bottom:1px solid #e2e8f0;">Faltante</th>
                        <th align="left" style="padding:9px;border-bottom:1px solid #e2e8f0;">Unidad</th>
                    </tr>
                </thead>
                <tbody>{''.join(filas)}</tbody>
            </table>
        </div>
        """

    texto = (
        f"{titulo}\n\n"
        f"PAW #{paw_numero} - {paw_nombre}\n"
        f"Recepción actual: {porcentaje}%\n\n"
        f"{mensaje_estado}"
        f"{pendientes_texto}\n"
        f"Revisar recepción: {inventario_url}\n"
        f"Revisar Taller: {taller_url}\n"
    )

    html = f"""
    <!doctype html>
    <html>
    <body style="margin:0;padding:0;background:#eef2f7;font-family:Arial,Helvetica,sans-serif;color:#0f172a;">
        <table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"
               style="background:#eef2f7;padding:24px 12px;">
            <tr>
                <td align="center">
                    <table role="presentation" width="760" cellspacing="0" cellpadding="0" border="0"
                           style="width:760px;max-width:760px;background:#ffffff;border-radius:14px;overflow:hidden;
                                  box-shadow:0 4px 14px rgba(15,23,42,.08);">

                        <tr>
                            <td style="background:#0f172a;padding:20px 24px;">
                                <div style="font-size:12px;letter-spacing:.8px;font-weight:800;color:#cbd5e1;">
                                    IMPETUS CONTROL
                                </div>
                                <div style="font-size:24px;line-height:1.25;font-weight:800;color:#ffffff;margin-top:6px;">
                                    {titulo}
                                </div>
                            </td>
                        </tr>

                        <tr>
                            <td style="padding:24px;">
                                <div style="font-size:14px;color:#64748b;margin-bottom:8px;">
                                    PAW #{paw_numero}
                                </div>

                                <div style="font-size:19px;font-weight:800;color:#0f172a;margin-bottom:18px;">
                                    {paw_nombre}
                                </div>

                                <table role="presentation" width="100%" cellspacing="0" cellpadding="0" border="0"
                                       style="background:#f8fafc;border:1px solid #e2e8f0;border-radius:12px;">
                                    <tr>
                                        <td style="padding:18px 20px;">
                                            <div style="font-size:12px;font-weight:800;color:#64748b;
                                                        text-transform:uppercase;letter-spacing:.5px;">
                                                Avance de recepción
                                            </div>
                                            <div style="font-size:34px;line-height:1.1;font-weight:900;
                                                        color:{'#16a34a' if umbral == 100 else '#f59e0b'};
                                                        margin-top:6px;">
                                                {porcentaje}%
                                            </div>
                                            <div style="font-size:14px;line-height:1.6;color:#475569;margin-top:10px;">
                                                {mensaje_estado}
                                            </div>
                                        </td>
                                    </tr>
                                </table>

                                {pendientes_html}

                                <table role="presentation" cellspacing="0" cellpadding="0" border="0"
                                       style="margin-top:24px;">
                                    <tr>
                                        <td style="padding-right:10px;">
                                            <a href="{inventario_url}"
                                               style="display:inline-block;background:#2563eb;color:#ffffff;
                                                      text-decoration:none;padding:11px 16px;border-radius:8px;
                                                      font-size:13px;font-weight:800;">
                                                Revisar recepción
                                            </a>
                                        </td>
                                        <td>
                                            <a href="{taller_url}"
                                               style="display:inline-block;background:#16a34a;color:#ffffff;
                                                      text-decoration:none;padding:11px 16px;border-radius:8px;
                                                      font-size:13px;font-weight:800;">
                                                Revisar Taller
                                            </a>
                                        </td>
                                    </tr>
                                </table>

                                <div style="margin-top:24px;padding-top:16px;border-top:1px solid #e5e7eb;
                                            color:#94a3b8;font-size:11px;line-height:1.5;">
                                    Notificación automática generada por Impetus Control al alcanzar este umbral
                                    por primera vez.
                                </div>
                            </td>
                        </tr>

                    </table>
                </td>
            </tr>
        </table>
    </body>
    </html>
    """

    enviados = 0
    for correo in destinatarios:
        try:
            msg = EmailMultiAlternatives(
                subject=asunto,
                body=texto,
                from_email=settings.DEFAULT_FROM_EMAIL,
                to=[correo],
            )
            msg.attach_alternative(html, "text/html")
            enviados += msg.send(fail_silently=False)
        except Exception:
            # Una dirección inválida no debe impedir que Inventario guarde la recepción.
            continue

    return enviados


@login_required
@inventario_required
def recepcion_detail(request, pk):
    recepcion = get_object_or_404(
        InventoryReception.objects
        .select_related("purchase_request", "creado_por")
        .prefetch_related("lineas__purchase_line"),
        pk=pk
    )

    # Solo deben aparecer en Recepción los ítems que realmente quedaron para compra.
    # Esto también evita que líneas históricas con cantidad 0 afecten estados/progreso.
    lineas_recepcion = recepcion.lineas.filter(
        purchase_line__cantidad_a_comprar__gt=0,
        cantidad_esperada__gt=0,
    ).select_related("purchase_line")

    # Corrige recepciones antiguas que fueron creadas sin código/descripcion/unidad.
    for linea in lineas_recepcion:
        if linea.purchase_line:
            actualizado = False

            if not linea.codigo:
                linea.codigo = linea.purchase_line.codigo or ""
                actualizado = True

            if not linea.descripcion:
                linea.descripcion = linea.purchase_line.descripcion or ""
                actualizado = True

            if not linea.unidad:
                linea.unidad = linea.purchase_line.unidad or ""
                actualizado = True

            if actualizado:
                linea.save(update_fields=["codigo", "descripcion", "unidad"])

    if request.method == "POST":
        # La recepción ahora afecta existencias por el INCREMENTO recibido.
        # Una compra originada por faltante del PAW entra a físico y queda
        # reservada inmediatamente para ese mismo PAW.
        from .models import InventoryStock, InventoryMovement, InventoryReservation

        errores_recepcion = []
        datos_recepcion = []

        for linea in lineas_recepcion:
            raw = request.POST.get(f"cantidad_recibida_{linea.id}") or "0"
            try:
                cantidad = Decimal(raw.replace(",", "."))
            except Exception:
                errores_recepcion.append(f"Cantidad inválida para {linea.codigo or linea.descripcion}.")
                continue

            anterior = Decimal(linea.cantidad_recibida or 0)
            esperada = Decimal(linea.cantidad_esperada or 0)

            if cantidad < 0 or cantidad > esperada:
                errores_recepcion.append(
                    f"{linea.codigo or linea.descripcion}: la cantidad recibida debe estar entre 0 y {esperada}."
                )
                continue

            # Una recepción ya contabilizada no se reduce editando la pantalla:
            # una corrección física debe hacerse mediante Ajuste de inventario.
            if cantidad < anterior:
                errores_recepcion.append(
                    f"{linea.codigo or linea.descripcion}: ya hay {anterior} recibidas. "
                    "No reduzcas una recepción contabilizada; usa Ajuste de inventario si necesitas corregir el físico."
                )
                continue

            datos_recepcion.append((
                linea, cantidad, anterior,
                request.POST.get(f"fecha_llegada_{linea.id}") or None,
                request.POST.get(f"observacion_{linea.id}") or "",
            ))

        if errores_recepcion:
            for error in errores_recepcion:
                messages.error(request, error)
            return redirect("inventario:recepcion_detail", pk=recepcion.pk)

        with transaction.atomic():
            for linea, cantidad, anterior, fecha, observacion in datos_recepcion:
                incremento = cantidad - anterior
                esperada = Decimal(linea.cantidad_esperada or 0)

                # Solo el incremento nuevo genera entrada física y Kardex.
                if incremento > 0:
                    codigo = (linea.codigo or getattr(linea.purchase_line, "codigo", "") or "").strip()
                    catalogo, CatalogModel = _catalog_model_for_empresa(InventoryStock.Empresa.IMPETUS)
                    item = CatalogModel.objects.filter(codigo__iexact=codigo).first() if codigo else None
                    if not item:
                        transaction.set_rollback(True)
                        messages.error(
                            request,
                            f"No se encontró el P/N {codigo or '-'} en el catálogo IMPETUS. "
                            "No se contabilizó la recepción."
                        )
                        return redirect("inventario:recepcion_detail", pk=recepcion.pk)

                    stock = _get_or_create_stock(InventoryStock.Empresa.IMPETUS, item)
                    stock = InventoryStock.objects.select_for_update().get(pk=stock.pk)

                    saldo_anterior = Decimal(stock.cantidad_fisica or 0)
                    costo_anterior = Decimal(stock.costo_promedio or 0)
                    precio_compra = Decimal(getattr(linea.purchase_line, "precio_unitario", 0) or 0)

                    saldo_nuevo = saldo_anterior + incremento
                    if precio_compra > 0 and saldo_nuevo > 0:
                        costo_nuevo = ((saldo_anterior * costo_anterior) + (incremento * precio_compra)) / saldo_nuevo
                    else:
                        costo_nuevo = costo_anterior

                    stock.cantidad_fisica = saldo_nuevo
                    stock.costo_promedio = costo_nuevo

                    # Reserva automática: como esta línea nació del faltante de este PAW,
                    # lo recibido no queda libre para otro PAW.
                    requerida = Decimal(getattr(linea.purchase_line, "cantidad_requerida", 0) or 0)
                    ya_reservado_linea = InventoryReservation.objects.filter(
                        purchase_line=linea.purchase_line,
                        estado=InventoryReservation.Estado.ACTIVA,
                    ).aggregate(total=Sum("cantidad"))["total"] or Decimal("0")
                    pendiente_reserva = max(requerida - Decimal(ya_reservado_linea), Decimal("0"))
                    a_reservar = min(incremento, pendiente_reserva)

                    if a_reservar > 0:
                        InventoryReservation.objects.create(
                            stock=stock,
                            cantidad=a_reservar,
                            purchase_request=recepcion.purchase_request,
                            purchase_line=linea.purchase_line,
                            creado_por=request.user,
                            observacion=f"Reserva automática por recepción PAW #{recepcion.purchase_request.paw_numero}",
                        )
                        stock.cantidad_reservada = Decimal(stock.cantidad_reservada or 0) + a_reservar

                    stock.save(update_fields=["cantidad_fisica", "cantidad_reservada", "costo_promedio", "actualizado_en"])

                    InventoryMovement.objects.create(
                        stock=stock,
                        tipo=InventoryMovement.Tipo.RECEPCION,
                        cantidad=incremento,
                        costo_unitario=precio_compra if precio_compra > 0 else costo_nuevo,
                        saldo_anterior=saldo_anterior,
                        saldo_nuevo=saldo_nuevo,
                        costo_promedio_anterior=costo_anterior,
                        costo_promedio_nuevo=costo_nuevo,
                        referencia=f"PAW-{recepcion.purchase_request.paw_numero}",
                        motivo=f"Recepción de compra. Reserva automática: {a_reservar}",
                        creado_por=request.user,
                    )

                linea.cantidad_recibida = cantidad
                linea.fecha_llegada = fecha
                linea.observacion_inventario = observacion
                if cantidad <= 0:
                    linea.estado = "PENDIENTE"
                elif cantidad < esperada:
                    linea.estado = "PARCIAL"
                else:
                    linea.estado = "LISTO"
                linea.save()

        total = lineas_recepcion.count()
        listas = lineas_recepcion.filter(estado="LISTO").count()
        parciales = lineas_recepcion.filter(estado="PARCIAL").count()

        try:
            paw = recepcion.purchase_request.bom.workorder.paw

            if total > 0 and listas == total:
                paw.estado_operativo = "MATERIAL_RECIBIDO"
            elif listas > 0 or parciales > 0:
                paw.estado_operativo = "MATERIAL_PARCIAL"

            paw.save(update_fields=["estado_operativo"])
        except Exception:
            pass

        # ======================================================
        # ALERTAS EN VIVO DE RECEPCIÓN: 80% Y 100%
        # Se calcula por cantidad recibida / cantidad esperada,
        # no por número de líneas, para representar mejor el avance real.
        # ======================================================
        cantidad_esperada_total = Decimal("0")
        cantidad_recibida_total = Decimal("0")
        pendientes = []

        for linea in lineas_recepcion:
            esperada = Decimal(linea.cantidad_esperada or 0)
            recibida = Decimal(linea.cantidad_recibida or 0)

            if esperada <= 0:
                continue

            cantidad_esperada_total += esperada
            cantidad_recibida_total += min(max(recibida, Decimal("0")), esperada)

            faltante = max(esperada - recibida, Decimal("0"))
            if faltante > 0:
                pendientes.append({
                    "codigo": linea.codigo or "-",
                    "descripcion": linea.descripcion or "",
                    "faltante": f"{faltante.normalize()}",
                    "unidad": linea.unidad or "",
                })

        if cantidad_esperada_total > 0:
            porcentaje_recepcion = int(
                (cantidad_recibida_total / cantidad_esperada_total) * Decimal("100")
            )
            porcentaje_recepcion = min(100, max(0, porcentaje_recepcion))
        else:
            porcentaje_recepcion = 0

        # Primero 80%. Si una recepción pasa directamente de <80 a 100,
        # se envía únicamente la alerta de 100% para evitar dos correos simultáneos.
        if porcentaje_recepcion >= 100 and not recepcion.notificacion_100_en:
            enviados = _enviar_alerta_recepcion(
                recepcion=recepcion,
                porcentaje=100,
                pendientes=[],
                umbral=100,
            )
            if enviados:
                ahora = timezone.now()
                recepcion.notificacion_100_en = ahora
                # También marcamos 80 como cumplido: no debe enviarse después.
                if not recepcion.notificacion_80_en:
                    recepcion.notificacion_80_en = ahora
                recepcion.save(
                    update_fields=[
                        "notificacion_80_en",
                        "notificacion_100_en",
                        "actualizado_en",
                    ]
                )

        elif porcentaje_recepcion >= 80 and not recepcion.notificacion_80_en:
            enviados = _enviar_alerta_recepcion(
                recepcion=recepcion,
                porcentaje=porcentaje_recepcion,
                pendientes=pendientes,
                umbral=80,
            )
            if enviados:
                recepcion.notificacion_80_en = timezone.now()
                recepcion.save(
                    update_fields=["notificacion_80_en", "actualizado_en"]
                )

        messages.success(request, "Recepción de inventario actualizada correctamente.")
        return redirect("inventario:recepcion_detail", pk=recepcion.pk)

    return render(request, "inventario/recepcion_detail.html", {
        "recepcion": recepcion,
        "lineas": lineas_recepcion,
    })


@login_required
@inventario_required
def entrega_taller_detail(request, pk):
    """Registra la entrega física y consume exclusivamente reservas del mismo PAW/línea."""
    from .models import InventoryReservation, InventoryMovement, InventoryStock

    entrega = get_object_or_404(
        WorkshopDelivery.objects
        .select_related("purchase_request", "creado_por")
        .prefetch_related("lineas__purchase_line"),
        pk=pk
    )

    # Información de reserva pendiente por línea para mostrarla en pantalla.
    def _reserva_pendiente(purchase_line_id):
        reservas = InventoryReservation.objects.filter(
            purchase_line_id=purchase_line_id,
            purchase_request=entrega.purchase_request,
            estado=InventoryReservation.Estado.ACTIVA,
        )
        total = Decimal("0")
        for r in reservas:
            total += max(Decimal(r.cantidad or 0) - Decimal(r.cantidad_consumida or 0), Decimal("0"))
        return total

    if request.method == "POST":
        errores = []
        datos = []

        # Validación completa ANTES de descontar cualquier existencia.
        for linea in entrega.lineas.all():
            requerida = Decimal(linea.cantidad_requerida or 0)
            anterior = Decimal(linea.cantidad_entregada or 0)
            raw = request.POST.get(f"cantidad_entregada_{linea.id}")

            # Inputs deshabilitados (líneas completas) no llegan en POST.
            if raw is None or raw == "":
                nueva = anterior
            else:
                try:
                    nueva = Decimal(str(raw).replace(",", "."))
                except Exception:
                    errores.append(f"Cantidad inválida para {linea.codigo or linea.descripcion}.")
                    continue

            if nueva < anterior:
                errores.append(
                    f"{linea.codigo or linea.descripcion}: no puedes reducir una entrega ya contabilizada "
                    f"({anterior}). Para corregir una salida debe hacerse un movimiento de devolución/ajuste."
                )
                continue
            if nueva > requerida:
                errores.append(
                    f"{linea.codigo or linea.descripcion}: no puedes entregar {nueva}; el requerido es {requerida}."
                )
                continue

            incremento = nueva - anterior
            if incremento > 0:
                reserva = _reserva_pendiente(linea.purchase_line_id)
                if incremento > reserva:
                    errores.append(
                        f"{linea.codigo or linea.descripcion}: intentas entregar {incremento} adicional, "
                        f"pero este PAW solo tiene {reserva} reservado pendiente."
                    )
                    continue
            datos.append((linea, anterior, nueva, incremento))

        if errores:
            for error in errores:
                messages.error(request, error)
        else:
            try:
                with transaction.atomic():
                    for linea, anterior, nueva, incremento in datos:
                        if incremento <= 0:
                            continue

                        # Bloqueamos y consumimos SOLO reservas de esta PurchaseLine/PAW.
                        reservas = list(
                            InventoryReservation.objects.select_for_update()
                            .select_related("stock")
                            .filter(
                                purchase_line=linea.purchase_line,
                                purchase_request=entrega.purchase_request,
                                estado=InventoryReservation.Estado.ACTIVA,
                            )
                            .order_by("creado_en", "id")
                        )
                        pendiente = incremento

                        # Agrupamos por stock por seguridad (normalmente será un único P/N/stock).
                        consumo_por_stock = {}
                        for reserva in reservas:
                            disponible_reserva = max(
                                Decimal(reserva.cantidad or 0) - Decimal(reserva.cantidad_consumida or 0),
                                Decimal("0"),
                            )
                            tomar = min(disponible_reserva, pendiente)
                            if tomar <= 0:
                                continue
                            reserva.cantidad_consumida = Decimal(reserva.cantidad_consumida or 0) + tomar
                            if reserva.cantidad_consumida >= Decimal(reserva.cantidad or 0):
                                reserva.estado = (
                                    InventoryReservation.Estado.LIBERADA
                                    if entrega.destino == "INVENTARIO"
                                    else InventoryReservation.Estado.CONSUMIDA
                                )
                                reserva.cerrado_en = timezone.now()
                                reserva.save(update_fields=["cantidad_consumida", "estado", "cerrado_en"])
                            else:
                                reserva.save(update_fields=["cantidad_consumida"])
                            consumo_por_stock[reserva.stock_id] = consumo_por_stock.get(reserva.stock_id, Decimal("0")) + tomar
                            pendiente -= tomar
                            if pendiente <= 0:
                                break

                        if pendiente > 0:
                            raise ValueError(
                                f"La reserva de {linea.codigo or linea.descripcion} cambió mientras se procesaba la entrega."
                            )

                        for stock_id, consumo in consumo_por_stock.items():
                            stock = InventoryStock.objects.select_for_update().get(pk=stock_id)
                            fisico_anterior = Decimal(stock.cantidad_fisica or 0)
                            reservado_anterior = Decimal(stock.cantidad_reservada or 0)
                            if consumo > reservado_anterior:
                                raise ValueError(
                                    f"{stock.codigo}: inconsistencia de reserva. Reservado {reservado_anterior}, proceso {consumo}."
                                )

                            if entrega.destino == "INVENTARIO":
                                # Compra para stock/bodega: la recepción ya aumentó el físico.
                                # Aquí NO hay salida física; únicamente liberamos la reserva
                                # del PAW para que las unidades queden disponibles.
                                stock.cantidad_reservada = reservado_anterior - consumo
                                stock.save(update_fields=["cantidad_reservada", "actualizado_en"])
                            else:
                                if consumo > fisico_anterior:
                                    raise ValueError(
                                        f"{stock.codigo}: existencia física insuficiente. Físico {fisico_anterior}, salida {consumo}."
                                    )
                                stock.cantidad_fisica = fisico_anterior - consumo
                                stock.cantidad_reservada = reservado_anterior - consumo
                                stock.save(update_fields=["cantidad_fisica", "cantidad_reservada", "actualizado_en"])

                                InventoryMovement.objects.create(
                                    stock=stock,
                                    tipo=InventoryMovement.Tipo.ENTREGA,
                                    cantidad=-consumo,
                                    costo_unitario=Decimal(stock.costo_promedio or 0),
                                    saldo_anterior=fisico_anterior,
                                    saldo_nuevo=Decimal(stock.cantidad_fisica or 0),
                                    costo_promedio_anterior=Decimal(stock.costo_promedio or 0),
                                    costo_promedio_nuevo=Decimal(stock.costo_promedio or 0),
                                    referencia=f"PAW-{entrega.purchase_request.paw_numero}",
                                    motivo=(
                                        f"Entrega a {entrega.get_destino_display()} - consumo de reserva del PAW. "
                                        f"P/N {linea.codigo or '-'}"
                                    ),
                                    creado_por=request.user,
                                )

                        linea.cantidad_entregada = nueva
                        linea.save(update_fields=["cantidad_entregada"])

                    entrega.comentarios = request.POST.get("comentarios", "")
                    entrega.save(update_fields=["comentarios", "actualizado_en"])

                    completa = True
                    for linea in entrega.lineas.all():
                        req = Decimal(linea.cantidad_requerida or 0)
                        ent = Decimal(linea.cantidad_entregada or 0)
                        if req > 0 and ent < req:
                            completa = False
                            break

                    if completa:
                        try:
                            paw = entrega.purchase_request.bom.workorder.paw
                            destino = getattr(entrega, "destino", "TALLER")
                            if destino == "TALLER" and getattr(paw, "aplica_taller", True):
                                paw.estado_operativo = "ENTREGADO_TALLER"
                            elif destino == "INVENTARIO" and not getattr(paw, "aplica_taller", False) and not getattr(paw, "aplica_campo", False):
                                paw.estado_operativo = "PRODUCTO_OK"
                            else:
                                paw.estado_operativo = "MATERIAL_RECIBIDO"
                            paw.save(update_fields=["estado_operativo"])
                        except Exception:
                            pass

                    registrar_movimiento(
                        request=request,
                        paw_numero=entrega.purchase_request.paw_numero,
                        modulo="INVENTARIO",
                        accion=f"Entrega a {entrega.get_destino_display()}",
                        descripcion=("Ingreso a bodega y liberación de reservas del PAW." if entrega.destino == "INVENTARIO" else "Salida física de inventario y consumo de reservas del PAW."),
                        objeto=entrega,
                    )

                messages.success(
                    request,
                    (f"Ingreso a {entrega.get_destino_display()} confirmado. El stock físico permanece y la reserva del PAW fue liberada." if entrega.destino == "INVENTARIO" else f"Entrega a {entrega.get_destino_display()} actualizada. Stock físico y reservas descontados correctamente.")
                )
                return redirect("inventario:entrega_taller_detail", pk=entrega.pk)
            except ValueError as exc:
                messages.error(request, str(exc))

    # Valores informativos para el template.
    for linea in entrega.lineas.all():
        linea.reserva_pendiente_paw = _reserva_pendiente(linea.purchase_line_id)

    return render(request, "inventario/entrega_taller_detail.html", {"entrega": entrega})


@login_required
@inventario_required
def entrega_taller_pdf(request, pk):
    entrega = get_object_or_404(
        WorkshopDelivery.objects
        .select_related("purchase_request")
        .prefetch_related("lineas"),
        pk=pk
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28,
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"<b>ENTREGA {entrega.get_destino_display().upper()}</b>", styles["Title"]))
    story.append(Spacer(1, 8))

    story.append(Paragraph(
        f"<b>Paw #:</b> {entrega.purchase_request.paw_numero} "
        f"&nbsp;&nbsp;&nbsp; <b>Nombre PAW:</b> {entrega.purchase_request.paw_nombre}",
        styles["Normal"]
    ))

    story.append(Paragraph(
        f"<b>Fecha impresión:</b> {timezone.now().date()}",
        styles["Normal"]
    ))

    story.append(Spacer(1, 12))

    data = [[
        "CÓDIGO",
        "DESCRIPCIÓN",
        "UNID",
        "CANT. REQ",
        "CANT. ENT",
    ]]

    for linea in entrega.lineas.all():
        data.append([
            linea.codigo or "",
            Paragraph(linea.descripcion or "", styles["Normal"]),
            linea.unidad or "",
            f"{Decimal(linea.cantidad_requerida or 0):.0f}",
            f"{Decimal(linea.cantidad_entregada or 0):.0f}",
        ])

    table = Table(
        data,
        colWidths=[70, 270, 45, 65, 65],
        repeatRows=1,
    )

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    story.append(table)

    story.append(Spacer(1, 18))
    story.append(Paragraph("<b>Comentarios</b>", styles["Heading3"]))
    story.append(Paragraph(entrega.comentarios or " ", styles["Normal"]))

    story.append(Spacer(1, 36))
    story.append(Paragraph("<b>Firmas</b>", styles["Heading3"]))
    story.append(Spacer(1, 32))

    firmas = Table([
        ["__________________________", "__________________________"],
        ["Firma entrega", "Firma recibe (Taller)"],
        ["", ""],
        ["Fecha", "Fecha"],
    ], colWidths=[250, 250])

    firmas.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))

    story.append(firmas)
    doc.build(story)

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'inline; filename="ENTREGA_TALLER_{entrega.purchase_request.paw_numero}.pdf"'
    )
    return response
# ======================================================
# SALIDAS DE INVENTARIO SIN PAW
# ======================================================

def _decimal_post(value, default="0"):
    try:
        return Decimal(str(value or default).replace(",", "."))
    except Exception:
        return Decimal(default)


def _lineas_desde_post(request, prefix):
    """Lee filas dinámicas enviadas como prefix_descripcion[], etc."""
    descripciones = request.POST.getlist(f"{prefix}_descripcion[]")
    codigos = request.POST.getlist(f"{prefix}_codigo[]")
    cantidades = request.POST.getlist(f"{prefix}_cantidad[]")
    unidades = request.POST.getlist(f"{prefix}_unidad[]")
    seriales = request.POST.getlist(f"{prefix}_serial[]")
    catalogos = request.POST.getlist(f"{prefix}_catalogo[]")
    catalogo_ids = request.POST.getlist(f"{prefix}_catalogo_id[]")

    total = max(len(descripciones), len(codigos), len(cantidades), 0)
    filas = []
    for i in range(total):
        descripcion = (descripciones[i] if i < len(descripciones) else "").strip()
        codigo = (codigos[i] if i < len(codigos) else "").strip()
        if not descripcion and not codigo:
            continue
        filas.append({
            "descripcion": descripcion or codigo,
            "codigo": codigo,
            "cantidad": max(_decimal_post(cantidades[i] if i < len(cantidades) else "1", "1"), Decimal("0")),
            "unidad": (unidades[i] if i < len(unidades) else "").strip(),
            "serial": (seriales[i] if i < len(seriales) else "").strip(),
            "catalogo": (catalogos[i] if i < len(catalogos) else "").strip(),
            "catalogo_id": int(catalogo_ids[i]) if i < len(catalogo_ids) and str(catalogo_ids[i]).isdigit() else None,
        })
    return filas


@login_required
@inventario_required
def salidas_lista(request):
    salidas = InventoryExit.objects.select_related("creado_por").prefetch_related("lineas").order_by("-creado_en")
    return render(request, "inventario/salidas_lista.html", {"salidas": salidas})


@login_required
@inventario_required
def salida_nueva(request):
    if request.method == "POST":
        filas = _lineas_desde_post(request, "item")
        if not filas:
            messages.error(request, "Agrega al menos un componente a la salida.")
        else:
            with transaction.atomic():
                salida = InventoryExit.objects.create(
                    destino=(request.POST.get("destino") or "").strip(),
                    solicitado_por=(request.POST.get("solicitado_por") or "").strip(),
                    recibido_por=(request.POST.get("recibido_por") or "").strip(),
                    motivo=(request.POST.get("motivo") or "").strip(),
                    comentarios=(request.POST.get("comentarios") or "").strip(),
                    creado_por=request.user,
                )
                InventoryExitLine.objects.bulk_create([
                    InventoryExitLine(
                        salida=salida,
                        catalogo=f["catalogo"],
                        catalogo_item_id=f["catalogo_id"],
                        codigo=f["codigo"],
                        descripcion=f["descripcion"],
                        unidad=f["unidad"],
                        cantidad=f["cantidad"],
                        numero_serial=f["serial"],
                    ) for f in filas
                ])
            messages.success(request, f"Salida {salida.codigo} creada correctamente.")
            return redirect("inventario:salida_detail", pk=salida.pk)

    return render(request, "inventario/salida_form.html")


@login_required
@inventario_required
def salida_detail(request, pk):
    salida = get_object_or_404(InventoryExit.objects.select_related("creado_por").prefetch_related("lineas"), pk=pk)
    return render(request, "inventario/salida_detail.html", {"salida": salida})


@login_required
@inventario_required
def salida_pdf(request, pk):
    salida = get_object_or_404(InventoryExit.objects.select_related("creado_por").prefetch_related("lineas"), pk=pk)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>SALIDA DE INVENTARIO</b>", styles["Title"]),
        Spacer(1, 8),
        Paragraph(f"<b>Código:</b> {salida.codigo} &nbsp;&nbsp; <b>Fecha:</b> {timezone.localtime(salida.creado_en).strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
        Paragraph(f"<b>Destino:</b> {salida.destino or '-'}", styles["Normal"]),
        Paragraph(f"<b>Solicitado por:</b> {salida.solicitado_por or '-'} &nbsp;&nbsp; <b>Recibido por:</b> {salida.recibido_por or '-'}", styles["Normal"]),
        Paragraph(f"<b>Motivo:</b> {salida.motivo or '-'}", styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [["CÓDIGO / P.N.", "DESCRIPCIÓN", "UNID", "CANT.", "SERIAL"]]
    for linea in salida.lineas.all():
        data.append([
            linea.codigo or "",
            Paragraph(linea.descripcion or "", styles["Normal"]),
            linea.unidad or "",
            str(linea.cantidad.normalize()),
            linea.numero_serial or "",
        ])
    table = Table(data, colWidths=[90, 255, 50, 55, 85], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
        ("GRID", (0,0), (-1,-1), 0.6, colors.grey),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 8.5),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (2,1), (-1,-1), "CENTER"),
        ("LEFTPADDING", (0,0), (-1,-1), 4),
        ("RIGHTPADDING", (0,0), (-1,-1), 4),
    ]))
    story += [table, Spacer(1, 14), Paragraph("<b>Comentarios</b>", styles["Heading3"]), Paragraph(salida.comentarios or " ", styles["Normal"]), Spacer(1, 38)]
    firmas = Table([
        ["__________________________", "__________________________"],
        ["Entrega Inventario", "Recibe"],
        [salida.creado_por.get_full_name() if salida.creado_por and salida.creado_por.get_full_name() else (salida.creado_por.username if salida.creado_por else ""), salida.recibido_por or ""],
    ], colWidths=[250,250])
    firmas.setStyle(TableStyle([("ALIGN", (0,0), (-1,-1), "CENTER"), ("FONTSIZE", (0,0), (-1,-1), 9)]))
    story.append(firmas)
    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="{salida.codigo}.pdf"'
    return response


# ======================================================
# REMISIONES DE SALIDA - F-IN-04
# ======================================================

@login_required
@inventario_required
def remisiones_lista(request):
    remisiones = DispatchRemission.objects.select_related("creado_por").prefetch_related("lineas").order_by("-creado_en")
    return render(request, "inventario/remisiones_lista.html", {"remisiones": remisiones})


@login_required
@inventario_required
def remision_nueva(request):
    from quotes.models import Cliente
    clientes = Cliente.objects.filter(activo=True).order_by("nombre")

    if request.method == "POST":
        filas = _lineas_desde_post(request, "item")
        cliente_id = request.POST.get("cliente_registrado")
        cliente_obj = Cliente.objects.filter(pk=cliente_id, activo=True).first()
        empresa = request.POST.get("empresa")
        empresas_validas = {DispatchRemission.Empresa.IMPETUS, DispatchRemission.Empresa.OIL_GAS}

        if not cliente_obj:
            messages.error(request, "Selecciona un cliente registrado.")
        elif empresa not in empresas_validas:
            messages.error(request, "Selecciona la empresa que emite la remisión.")
        elif not filas:
            messages.error(request, "Agrega al menos un ítem a la remisión.")
        else:
            # Seguridad: un ítem seleccionado del catálogo debe pertenecer a la
            # misma empresa que emite la remisión. Los ítems diligenciados
            # manualmente pueden venir sin catálogo.
            catalogo_esperado = (
                "IMPETUS"
                if empresa == DispatchRemission.Empresa.IMPETUS
                else "OIL_GAS"
            )
            filas_invalidas = [
                f for f in filas
                if f["catalogo"] and f["catalogo"] != catalogo_esperado
            ]
            if filas_invalidas:
                messages.error(
                    request,
                    "Hay ítems seleccionados de un catálogo que no corresponde "
                    "a la empresa emisora. Vuelve a seleccionarlos."
                )
                return render(request, "inventario/remision_form.html", {
                    "fecha_hoy": timezone.localdate(),
                    "clientes": clientes,
                })

            with transaction.atomic():
                # Cada empresa conserva su propia serie documental.
                # OIL & GAS: última histórica OGS-RM-1511 -> inicia OGS-RM-1512.
                # IMPETUS:   última histórica REM-069     -> inicia REM-070.
                inicio = 69 if empresa == DispatchRemission.Empresa.IMPETUS else 1511
                secuencia, _ = RemissionSequence.objects.select_for_update().get_or_create(
                    empresa=empresa,
                    defaults={"ultimo": inicio},
                )
                secuencia.ultimo += 1
                secuencia.save(update_fields=["ultimo"])
                remision = DispatchRemission.objects.create(
                    consecutivo=secuencia.ultimo,
                    empresa=empresa,
                    cliente_registrado=cliente_obj,
                    # Se guarda una copia del nombre/NIT para conservar el histórico
                    # aunque posteriormente cambie el maestro de clientes.
                    cliente=cliente_obj.nombre,
                    nit=cliente_obj.nit or "",
                    fecha_envio=request.POST.get("fecha_envio") or timezone.localdate(),
                    contacto_envio=(request.POST.get("contacto_envio") or "").strip(),
                    telefono_envio=(request.POST.get("telefono_envio") or "").strip(),
                    direccion_envio=(request.POST.get("direccion_envio") or "").strip(),
                    tipo_vehiculo=(request.POST.get("tipo_vehiculo") or "").strip(),
                    placa=(request.POST.get("placa") or "").strip(),
                    nombre_conductor=(request.POST.get("nombre_conductor") or "").strip(),
                    celular_conductor=(request.POST.get("celular_conductor") or "").strip(),
                    observaciones=(request.POST.get("observaciones") or "").strip(),
                    creado_por=request.user,
                )
                DispatchRemissionLine.objects.bulk_create([
                    DispatchRemissionLine(
                        remision=remision, catalogo=f["catalogo"], catalogo_item_id=f["catalogo_id"],
                        descripcion=f["descripcion"], cantidad=f["cantidad"], unidad=f["unidad"] or "UND",
                        parte_numero=f["codigo"], numero_serial=f["serial"],
                    ) for f in filas
                ])
            messages.success(request, f"Remisión {remision.numero} creada correctamente.")
            return redirect("inventario:remision_detail", pk=remision.pk)

    return render(request, "inventario/remision_form.html", {
        "fecha_hoy": timezone.localdate(),
        "clientes": clientes,
    })


@login_required
@inventario_required
def remision_detail(request, pk):
    remision = get_object_or_404(
        DispatchRemission.objects.select_related("creado_por", "anulada_por").prefetch_related("lineas"),
        pk=pk,
    )
    return render(request, "inventario/remision_detail.html", {"remision": remision})


@login_required
@inventario_required
@require_POST
def remision_anular(request, pk):
    motivo = (request.POST.get("motivo_anulacion") or "").strip()
    if not motivo:
        messages.error(request, "Debes indicar el motivo de la anulación.")
        return redirect("inventario:remision_detail", pk=pk)

    with transaction.atomic():
        remision = get_object_or_404(
            DispatchRemission.objects.select_for_update(),
            pk=pk,
        )
        if remision.estado == DispatchRemission.Estado.ANULADA:
            messages.warning(request, f"La remisión {remision.numero} ya se encuentra anulada.")
            return redirect("inventario:remision_detail", pk=pk)

        remision.estado = DispatchRemission.Estado.ANULADA
        remision.anulada_por = request.user
        remision.anulada_en = timezone.now()
        remision.motivo_anulacion = motivo
        remision.save(update_fields=[
            "estado", "anulada_por", "anulada_en", "motivo_anulacion", "actualizado_en"
        ])

    messages.success(request, f"Remisión {remision.numero} anulada. Se conserva en el historial.")
    return redirect("inventario:remision_detail", pk=pk)


@login_required
@inventario_required
def buscar_items_inventario(request):
    """
    Autocomplete del inventario según la empresa emisora de la remisión.
    IMPETUS consulta únicamente ItemImpetus y OIL_GAS únicamente Item.
    """
    q = (request.GET.get("q") or "").strip()
    empresa = (request.GET.get("empresa") or "").strip().upper()

    if len(q) < 2:
        return JsonResponse({"results": []})

    from item_oil_gas.models import Item, ItemImpetus

    if empresa == DispatchRemission.Empresa.IMPETUS:
        catalogo = "IMPETUS"
        Model = ItemImpetus
    elif empresa == DispatchRemission.Empresa.OIL_GAS:
        catalogo = "OIL_GAS"
        Model = Item
    else:
        return JsonResponse({"results": []})

    qs = (
        Model.objects
        .filter(activo=True)
        .filter(Q(codigo__icontains=q) | Q(descripcion__icontains=q))
        .order_by("codigo")[:20]
    )

    resultados = [{
        "catalogo": catalogo,
        "id": item.pk,
        "codigo": item.codigo or "",
        "descripcion": item.descripcion or "",
        "unidad": item.unidad_medida or "UND",
        "label": f"{item.codigo} - {item.descripcion[:100]}",
    } for item in qs]

    return JsonResponse({"results": resultados})


def _p(text, styles, bold=False):
    value = str(text or "")
    if bold:
        value = f"<b>{value}</b>"
    return Paragraph(value, styles["Normal"])


@login_required
@inventario_required
def remision_pdf(request, pk):
    remision = get_object_or_404(
        DispatchRemission.objects.select_related("creado_por", "cliente_registrado", "anulada_por").prefetch_related("lineas"), pk=pk
    )
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=24, rightMargin=24, topMargin=22, bottomMargin=22)
    styles = getSampleStyleSheet()
    story = []

    # Logo dinámico según la empresa emisora.
    # finders.find funciona tanto en desarrollo como con STATIC_ROOT/collectstatic
    # en producción (Render/WhiteNoise).
    logo_name = (
        "img/logo_empresa.png"
        if remision.empresa == DispatchRemission.Empresa.IMPETUS
        else "img/logo_oil_gas.png"
    )
    logo_path = finders.find(logo_name)
    logo = ""
    if logo_path:
        logo = Image(str(logo_path), width=92, height=48, kind="proportional")

    titulo = f"REMISIÓN {remision.empresa_nombre}"
    encabezado = Table([
        [logo, _p(titulo, styles, True), _p("Versión: 4", styles)],
        ["", _p("Copia controlada (X)    Copia no controlada ( )", styles), _p("Fecha: 01/06/2022", styles)],
        ["", _p(f"NIT {remision.empresa_nit}", styles, True), _p("Código: F-IN-04", styles)],
    ], colWidths=[105, 300, 120], rowHeights=[28, 25, 25])
    encabezado.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.65, colors.black),
        ("SPAN", (0,0), (0,2)),
        ("ALIGN", (0,0), (0,2), "CENTER"),
        ("ALIGN", (1,0), (1,2), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("TOPPADDING", (0,0), (-1,-1), 3),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(encabezado)

    if remision.estado == DispatchRemission.Estado.ANULADA:
        aviso = Table([["REMISIÓN ANULADA"]], colWidths=[525], rowHeights=[28])
        aviso.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fee2e2")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#991b1b")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 13),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#991b1b")),
        ]))
        story += [Spacer(1, 5), aviso, Spacer(1, 5)]

    info = Table([
        [_p("Datos del Cliente", styles, True), "", _p("Información de Envío", styles, True), ""],
        [_p("Cliente", styles, True), _p(remision.cliente, styles), _p("REMISIÓN N°", styles, True), _p(remision.numero, styles, True)],
        [_p("NIT", styles, True), _p(remision.nit, styles), _p("Fecha Envío", styles, True), _p(remision.fecha_envio.strftime("%d/%m/%Y"), styles)],
        [_p("Contacto", styles, True), _p(remision.contacto_envio, styles), _p("Teléfono", styles, True), _p(remision.telefono_envio, styles)],
        [_p("Dirección", styles, True), _p(remision.direccion_envio, styles), _p("Nombre Conductor", styles, True), _p(remision.nombre_conductor, styles)],
        [_p("Tipo de Vehículo", styles, True), _p(remision.tipo_vehiculo, styles), _p("Celular", styles, True), _p(remision.celular_conductor, styles)],
        [_p("Placa", styles, True), _p(remision.placa, styles), "", ""],
    ], colWidths=[82, 180, 92, 171])
    info.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.65, colors.black),
        ("SPAN", (0,0), (1,0)), ("SPAN", (2,0), (3,0)),
        ("BACKGROUND", (0,0), (1,0), colors.HexColor("#e8eef5")),
        ("BACKGROUND", (2,0), (3,0), colors.HexColor("#e8eef5")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(info)

    data = [["Ítem", "Descripción", "Cant.", "U/M", "Parte Número", "Número Serial"]]
    for idx, linea in enumerate(remision.lineas.all(), start=1):
        data.append([str(idx), _p(linea.descripcion, styles), str(linea.cantidad.normalize()), linea.unidad or "", linea.parte_numero or "", linea.numero_serial or ""])
    while len(data) < 14:
        idx = len(data)
        data.append([str(idx), "", "", "", "", ""])

    items = Table(data, colWidths=[32, 250, 45, 43, 78, 77], repeatRows=1, rowHeights=[24] + [26]*(len(data)-1))
    items.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.65, colors.black), ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#e8eef5")),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("ALIGN", (0,0), (0,-1), "CENTER"),
        ("ALIGN", (2,0), (-1,-1), "CENTER"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTSIZE", (0,0), (-1,-1), 7.3), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(items)
    if remision.observaciones:
        story += [Spacer(1, 7), Paragraph(f"<b>Observaciones:</b> {remision.observaciones}", styles["Normal"])]

    if remision.estado == DispatchRemission.Estado.ANULADA:
        anulador = (
            remision.anulada_por.get_full_name()
            if remision.anulada_por and remision.anulada_por.get_full_name()
            else (remision.anulada_por.username if remision.anulada_por else "")
        )
        fecha_anulacion = timezone.localtime(remision.anulada_en).strftime("%d/%m/%Y %H:%M") if remision.anulada_en else ""
        story += [
            Spacer(1, 8),
            Paragraph(f"<b>Motivo de anulación:</b> {remision.motivo_anulacion}", styles["Normal"]),
            Paragraph(f"<b>Anulada por:</b> {anulador} &nbsp;&nbsp; <b>Fecha:</b> {fecha_anulacion}", styles["Normal"]),
        ]

    story += [Spacer(1, 34)]
    usuario = remision.creado_por.get_full_name() if remision.creado_por and remision.creado_por.get_full_name() else (remision.creado_por.username if remision.creado_por else "")
    firmas = Table([
        ["______________________", "______________________", "______________________"],
        ["Despachado por", "Transportado por", "Recibido por"],
        ["Inventario", "Conductor", "Cliente"],
        [usuario, remision.nombre_conductor or "", ""],
    ], colWidths=[175,175,175])
    firmas.setStyle(TableStyle([
        ("ALIGN", (0,0), (-1,-1), "CENTER"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTNAME", (0,1), (-1,2), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 7.5),
    ]))
    story.append(firmas)

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="{remision.numero}.pdf"'
    return response


# ======================================================
# EXISTENCIAS / KARDEX - ETAPA 1
# ======================================================
def _catalog_model_for_empresa(empresa):
    from item_oil_gas.models import Item, ItemImpetus
    if empresa == "IMPETUS":
        return "IMPETUS", ItemImpetus
    if empresa == "OIL_GAS":
        return "OIL_GAS", Item
    return None, None


def _get_or_create_stock(empresa, item):
    from .models import InventoryStock
    catalogo, _ = _catalog_model_for_empresa(empresa)
    stock, _ = InventoryStock.objects.get_or_create(
        empresa=empresa, catalogo=catalogo, catalogo_item_id=item.pk,
        defaults={"codigo": item.codigo or "", "descripcion": item.descripcion or "", "unidad": item.unidad_medida or "UND"},
    )
    # Refresca snapshot del catálogo, nunca cantidades/costos.
    cambios = []
    for campo, valor in (("codigo", item.codigo or ""), ("descripcion", item.descripcion or ""), ("unidad", item.unidad_medida or "UND")):
        if getattr(stock, campo) != valor:
            setattr(stock, campo, valor); cambios.append(campo)
    if cambios:
        stock.save(update_fields=cambios + ["actualizado_en"])
    return stock


@login_required
@inventario_required
def existencias_lista(request):
    from .models import InventoryStock
    empresa = (request.GET.get("empresa") or "IMPETUS").upper()
    if empresa not in {"IMPETUS", "OIL_GAS"}: empresa = "IMPETUS"
    q = (request.GET.get("q") or "").strip()
    qs = InventoryStock.objects.filter(empresa=empresa)
    if q: qs = qs.filter(Q(codigo__icontains=q) | Q(descripcion__icontains=q))
    qs = qs.order_by("codigo")
    total_valor = sum((s.valor_inventario for s in qs), Decimal("0"))
    return render(request, "inventario/existencias_lista.html", {"stocks": qs, "empresa": empresa, "q": q, "total_valor": total_valor})


@login_required
@inventario_required
def inventario_inicial(request):
    from .models import InventoryMovement
    empresa = (request.GET.get("empresa") or request.POST.get("empresa") or "IMPETUS").upper()
    catalogo, Model = _catalog_model_for_empresa(empresa)
    if not Model:
        messages.error(request, "Empresa no válida."); return redirect("inventario:existencias_lista")
    q = (request.GET.get("q") or "").strip()
    items = list(Model.objects.filter(activo=True).filter(
        Q(codigo__icontains=q) | Q(descripcion__icontains=q)
    ).order_by("codigo")[:100]) if q else list(Model.objects.filter(activo=True).order_by("codigo")[:100])

    # Costo histórico sugerido: SOLO coincidencia exacta por P/N (PurchaseLine.codigo).
    # No cruza descripciones ni catálogos. Como referencia histórica, consulta todas las
    # compras registradas que tengan ese mismo P/N, independientemente del PAW.
    from compras_oil.models import PurchaseLine
    codigos = [str(i.codigo or "").strip() for i in items if str(i.codigo or "").strip()]
    historico = {}
    if codigos:
        lineas_precio = (
            PurchaseLine.objects
            .filter(codigo__in=codigos, precio_unitario__isnull=False, precio_unitario__gt=0, cantidad_a_comprar__gt=0)
            .select_related("request")
            .order_by("request__creado_en", "pk")
        )
        for linea in lineas_precio:
            codigo = (linea.codigo or "").strip()
            qty = Decimal(linea.cantidad_a_comprar or 0)
            precio = Decimal(linea.precio_unitario or 0)
            if qty <= 0 or precio <= 0:
                continue
            h = historico.setdefault(codigo, {"cantidad": Decimal("0"), "valor": Decimal("0"), "compras": 0, "ultimo": Decimal("0")})
            h["cantidad"] += qty
            h["valor"] += qty * precio
            h["compras"] += 1
            h["ultimo"] = precio

    filas = []
    for item in items:
        codigo = (item.codigo or "").strip()
        h = historico.get(codigo)
        costo_sugerido = Decimal("0")
        if h and h["cantidad"] > 0:
            costo_sugerido = (h["valor"] / h["cantidad"]).quantize(Decimal("0.01"))
        filas.append({
            "item": item,
            "costo_sugerido": costo_sugerido,
            "ultimo_precio": h["ultimo"] if h else Decimal("0"),
            "compras_encontradas": h["compras"] if h else 0,
        })

    if request.method == "POST":
        item = get_object_or_404(Model, pk=request.POST.get("item_id"), activo=True)
        try:
            cantidad = Decimal((request.POST.get("cantidad") or "0").replace(",", "."))
            costo = Decimal((request.POST.get("costo_unitario") or "0").replace(",", "."))
        except Exception:
            messages.error(request, "Cantidad o costo inválido."); return redirect(f"{request.path}?empresa={empresa}")
        if cantidad < 0 or costo < 0:
            messages.error(request, "Cantidad y costo no pueden ser negativos."); return redirect(f"{request.path}?empresa={empresa}")
        with transaction.atomic():
            stock = _get_or_create_stock(empresa, item)
            stock = type(stock).objects.select_for_update().get(pk=stock.pk)
            if stock.movimientos.exists() or stock.cantidad_fisica != 0:
                messages.error(request, "Este ítem ya tiene movimientos. Use Ajuste de inventario, no Inventario inicial.")
                return redirect(f"{request.path}?empresa={empresa}")
            anterior = stock.cantidad_fisica
            stock.cantidad_fisica = cantidad; stock.costo_promedio = costo
            stock.save(update_fields=["cantidad_fisica", "costo_promedio", "actualizado_en"])
            InventoryMovement.objects.create(stock=stock, tipo=InventoryMovement.Tipo.INVENTARIO_INICIAL, cantidad=cantidad, costo_unitario=costo, saldo_anterior=anterior, saldo_nuevo=cantidad, costo_promedio_anterior=0, costo_promedio_nuevo=costo, referencia="INVENTARIO INICIAL", motivo=(request.POST.get("motivo") or "Inventario inicial").strip(), creado_por=request.user)
        messages.success(request, f"Inventario inicial registrado: {item.codigo} = {cantidad} {item.unidad_medida or 'UND'}.")
        return redirect(f"{request.path}?empresa={empresa}")
    return render(request, "inventario/inventario_inicial.html", {"empresa": empresa, "filas": filas, "q": q})


@login_required
@inventario_required
@require_POST
def ajustar_stock(request, pk):
    from .models import InventoryStock, InventoryMovement
    try: nueva = Decimal((request.POST.get("cantidad_fisica") or "").replace(",", "."))
    except Exception:
        messages.error(request, "Cantidad física inválida."); return redirect("inventario:existencias_lista")
    motivo = (request.POST.get("motivo") or "").strip()
    if nueva < 0 or not motivo:
        messages.error(request, "La cantidad no puede ser negativa y el motivo es obligatorio."); return redirect("inventario:existencias_lista")
    with transaction.atomic():
        stock = InventoryStock.objects.select_for_update().get(pk=pk)
        if nueva < stock.cantidad_reservada:
            messages.error(request, f"No puede bajar de {stock.cantidad_reservada}; esa cantidad está reservada.")
            return redirect(f"/inventario/existencias/?empresa={stock.empresa}")
        anterior = stock.cantidad_fisica; delta = nueva - anterior
        if delta == 0:
            messages.info(request, "No hubo diferencia de inventario."); return redirect(f"/inventario/existencias/?empresa={stock.empresa}")
        tipo = InventoryMovement.Tipo.AJUSTE_ENTRADA if delta > 0 else InventoryMovement.Tipo.AJUSTE_SALIDA
        stock.cantidad_fisica = nueva; stock.save(update_fields=["cantidad_fisica", "actualizado_en"])
        InventoryMovement.objects.create(stock=stock, tipo=tipo, cantidad=delta, costo_unitario=stock.costo_promedio, saldo_anterior=anterior, saldo_nuevo=nueva, costo_promedio_anterior=stock.costo_promedio, costo_promedio_nuevo=stock.costo_promedio, referencia="AJUSTE", motivo=motivo, creado_por=request.user)
    messages.success(request, f"Ajuste registrado para {stock.codigo}: {anterior} → {nueva}.")
    return redirect(f"/inventario/existencias/?empresa={stock.empresa}")


@login_required
@inventario_required
def kardex_stock(request, pk):
    from .models import InventoryStock
    stock = get_object_or_404(InventoryStock, pk=pk)
    return render(request, "inventario/kardex_stock.html", {"stock": stock, "movimientos": stock.movimientos.select_related("creado_por").all(), "reservas": stock.reservas.select_related("purchase_request", "creado_por").all()})


@login_required
@inventario_required
def transferencia_items_destino(request):
    """Busca P/N destino en el catálogo de la otra empresa sin asumir equivalencias."""
    from django.http import JsonResponse
    from django.db.models import Q, Case, When, Value, IntegerField

    empresa = (request.GET.get("empresa") or "").upper()
    q = (request.GET.get("q") or "").strip()
    origen_id = request.GET.get("origen_id")
    if empresa not in {"IMPETUS", "OIL_GAS"}:
        return JsonResponse({"results": []})

    catalogo, Model = _catalog_model_for_empresa(empresa)
    qs = Model.objects.filter(activo=True)

    # Si no escribieron búsqueda, usa como ayuda la descripción/P/N del origen.
    origen = None
    if origen_id:
        try:
            from .models import InventoryStock
            origen = InventoryStock.objects.filter(pk=origen_id).first()
        except Exception:
            origen = None

    texto = q
    if not texto and origen:
        texto = origen.descripcion or origen.codigo or ""

    if texto:
        # Búsqueda segura por P/N exacto/parcial y por palabras significativas de descripción.
        tokens = [t for t in texto.replace("-", " ").replace("/", " ").split() if len(t) >= 3][:6]
        filtro = Q(codigo__iexact=texto) | Q(codigo__icontains=texto) | Q(descripcion__icontains=texto)
        for token in tokens:
            filtro |= Q(codigo__icontains=token) | Q(descripcion__icontains=token)
        qs = qs.filter(filtro).annotate(
            _rank=Case(
                When(codigo__iexact=texto, then=Value(0)),
                When(codigo__icontains=texto, then=Value(1)),
                When(descripcion__icontains=texto, then=Value(2)),
                default=Value(3), output_field=IntegerField(),
            )
        ).order_by("_rank", "codigo")
    else:
        qs = qs.order_by("codigo")

    results = [{
        "id": item.pk,
        "codigo": item.codigo or "",
        "descripcion": item.descripcion or "",
        "unidad": item.unidad_medida or "UND",
    } for item in qs[:30]]
    return JsonResponse({"results": results})


@login_required
@inventario_required
def transferencia_nueva(request):
    from .models import InventoryStock, InventoryMovement, InventoryTransfer
    if request.method == "POST":
        origen_id = request.POST.get("stock_origen")
        destino_empresa = (request.POST.get("empresa_destino") or "").upper()
        destino_item_id = request.POST.get("item_destino")
        motivo = (request.POST.get("motivo") or "").strip()
        documento = (request.POST.get("documento") or "").strip()
        try:
            cantidad = Decimal((request.POST.get("cantidad") or "0").replace(",", "."))
        except Exception:
            cantidad = Decimal("0")
        if cantidad <= 0 or not motivo:
            messages.error(request, "Cantidad mayor a cero y motivo son obligatorios.")
            return redirect("inventario:transferencia_nueva")
        if not destino_item_id:
            messages.error(request, "Debe seleccionar el P/N destino en el catálogo de la empresa receptora.")
            return redirect("inventario:transferencia_nueva")

        with transaction.atomic():
            origen = InventoryStock.objects.select_for_update().get(pk=origen_id)
            if destino_empresa == origen.empresa or destino_empresa not in {"IMPETUS", "OIL_GAS"}:
                messages.error(request, "Seleccione la otra empresa como destino.")
                return redirect("inventario:transferencia_nueva")
            if cantidad > origen.cantidad_disponible:
                messages.error(request, f"Disponible real: {origen.cantidad_disponible}. No se pueden transferir unidades reservadas.")
                return redirect("inventario:transferencia_nueva")

            dest_catalogo, DestModel = _catalog_model_for_empresa(destino_empresa)
            dest_item = DestModel.objects.filter(pk=destino_item_id, activo=True).first()
            if not dest_item:
                messages.error(request, "El P/N destino seleccionado no existe o está inactivo en el catálogo destino.")
                return redirect("inventario:transferencia_nueva")

            destino = _get_or_create_stock(destino_empresa, dest_item)
            destino = InventoryStock.objects.select_for_update().get(pk=destino.pk)
            costo = origen.costo_promedio

            oa = origen.cantidad_fisica
            origen.cantidad_fisica -= cantidad
            origen.save(update_fields=["cantidad_fisica", "actualizado_en"])

            da = destino.cantidad_fisica
            ca = destino.costo_promedio
            nuevo_total = da + cantidad
            nuevo_costo = ((da * ca) + (cantidad * costo)) / nuevo_total if nuevo_total else Decimal("0")
            destino.cantidad_fisica = nuevo_total
            destino.costo_promedio = nuevo_costo
            destino.save(update_fields=["cantidad_fisica", "costo_promedio", "actualizado_en"])

            trf = InventoryTransfer.objects.create(
                empresa_origen=origen.empresa, empresa_destino=destino_empresa,
                stock_origen=origen, stock_destino=destino,
                cantidad=cantidad, costo_unitario=costo,
                motivo=motivo, documento=documento, creado_por=request.user,
            )
            detalle = f"{motivo} | {origen.codigo} → {destino.codigo}"
            InventoryMovement.objects.create(
                stock=origen, tipo=InventoryMovement.Tipo.TRANSFERENCIA_SALIDA,
                cantidad=-cantidad, costo_unitario=costo, saldo_anterior=oa,
                saldo_nuevo=origen.cantidad_fisica, costo_promedio_anterior=costo,
                costo_promedio_nuevo=costo, referencia=trf.numero, motivo=detalle,
                creado_por=request.user,
            )
            InventoryMovement.objects.create(
                stock=destino, tipo=InventoryMovement.Tipo.TRANSFERENCIA_ENTRADA,
                cantidad=cantidad, costo_unitario=costo, saldo_anterior=da,
                saldo_nuevo=destino.cantidad_fisica, costo_promedio_anterior=ca,
                costo_promedio_nuevo=nuevo_costo, referencia=trf.numero,
                motivo=detalle, creado_por=request.user,
            )
        messages.success(request, f"Transferencia {trf.numero} completada: {origen.codigo} → {destino.codigo}.")
        return redirect("inventario:transferencias_lista")

    stocks = InventoryStock.objects.filter(cantidad_fisica__gt=F("cantidad_reservada")).order_by("empresa", "codigo")
    return render(request, "inventario/transferencia_form.html", {"stocks": stocks})


@login_required
@inventario_required
def transferencias_lista(request):
    from .models import InventoryTransfer
    return render(request, "inventario/transferencias_lista.html", {"transferencias": InventoryTransfer.objects.select_related("stock_origen", "stock_destino", "creado_por").all().order_by("-creado_en")})

# ======================================================
# INVENTARIO INICIAL - CARGA MASIVA EXCEL
# ======================================================
def _costos_historicos_por_codigo(codigos):
    """Costo ponderado histórico por P/N exacto. No cruza descripciones ni catálogos."""
    from compras_oil.models import PurchaseLine
    codigos = [str(c or "").strip() for c in codigos if str(c or "").strip()]
    historico = {}
    if not codigos:
        return historico
    lineas = (
        PurchaseLine.objects
        .filter(codigo__in=codigos, precio_unitario__isnull=False,
                precio_unitario__gt=0, cantidad_a_comprar__gt=0)
        .select_related("request")
        .order_by("request__creado_en", "pk")
    )
    for linea in lineas:
        codigo = (linea.codigo or "").strip()
        qty = Decimal(linea.cantidad_a_comprar or 0)
        precio = Decimal(linea.precio_unitario or 0)
        if qty <= 0 or precio <= 0:
            continue
        h = historico.setdefault(codigo, {
            "cantidad": Decimal("0"), "valor": Decimal("0"),
            "compras": 0, "ultimo": Decimal("0")
        })
        h["cantidad"] += qty
        h["valor"] += qty * precio
        h["compras"] += 1
        h["ultimo"] = precio
    return historico


@login_required
@inventario_required
def inventario_inicial_plantilla(request):
    """Descarga XLSX precargado con el catálogo activo de la empresa seleccionada."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "Falta instalar openpyxl para generar archivos Excel.")
        return redirect("inventario:inventario_inicial")

    empresa = (request.GET.get("empresa") or "IMPETUS").upper()
    catalogo, Model = _catalog_model_for_empresa(empresa)
    if not Model:
        messages.error(request, "Empresa no válida.")
        return redirect("inventario:inventario_inicial")

    items = list(Model.objects.filter(activo=True).order_by("codigo"))
    historico = _costos_historicos_por_codigo([i.codigo for i in items])

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario inicial"
    headers = ["P/N", "Descripción", "Unidad", "Cantidad física", "Costo unitario sugerido", "Observación"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for item in items:
        codigo = (item.codigo or "").strip()
        h = historico.get(codigo)
        costo = ""
        if h and h["cantidad"] > 0:
            costo = float((h["valor"] / h["cantidad"]).quantize(Decimal("0.01")))
        ws.append([
            codigo,
            item.descripcion or "",
            item.unidad_medida or "UND",
            "",  # El usuario diligencia únicamente el conteo físico real.
            costo,
            "",
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [18, 55, 14, 18, 24, 40]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "0.000"
        row[4].number_format = '#,##0.00'

    info = wb.create_sheet("INSTRUCCIONES")
    instrucciones = [
        "PLANTILLA DE INVENTARIO INICIAL",
        f"Empresa: {'IMPETUS HPS' if empresa == 'IMPETUS' else 'OIL & GAS SUPPORT'}",
        "Diligencie únicamente la cantidad que existe físicamente al momento del conteo.",
        "La cantidad física se registra independientemente de que el material esté reservado para un PAW.",
        "No cambie el P/N. La descripción y unidad son informativas.",
        "El costo sugerido proviene del historial de compras por P/N exacto y puede corregirse si corresponde.",
        "Deje Cantidad física vacía para los artículos que no desea cargar.",
        "Los artículos que ya tengan movimientos de inventario serán rechazados y deberán corregirse mediante Ajuste.",
    ]
    for line in instrucciones:
        info.append([line])
    info.column_dimensions["A"].width = 115
    info["A1"].font = Font(bold=True, size=14)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    nombre = f"inventario_inicial_{empresa.lower()}.xlsx"
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return response


@login_required
@inventario_required
def inventario_inicial_masivo(request):
    """Valida un XLSX y, tras confirmación, registra el inventario inicial en bloque."""
    from .models import InventoryStock, InventoryMovement
    empresa = (request.GET.get("empresa") or request.POST.get("empresa") or "IMPETUS").upper()
    catalogo, Model = _catalog_model_for_empresa(empresa)
    if not Model:
        messages.error(request, "Empresa no válida.")
        return redirect("inventario:inventario_inicial")

    session_key = f"inventario_inicial_masivo_{empresa}"
    contexto = {"empresa": empresa, "preview": None}

    if request.method == "POST" and request.POST.get("accion") == "cancelar":
        request.session.pop(session_key, None)
        messages.info(request, "Carga masiva cancelada. No se modificó el inventario.")
        return redirect(f"{request.path}?empresa={empresa}")

    if request.method == "POST" and request.POST.get("accion") == "confirmar":
        datos = request.session.get(session_key)
        if not datos:
            messages.error(request, "La validación expiró. Cargue nuevamente el archivo.")
            return redirect(f"{request.path}?empresa={empresa}")
        if datos.get("errores"):
            messages.error(request, "La carga contiene errores y no puede confirmarse.")
            return redirect(f"{request.path}?empresa={empresa}")

        filas = datos.get("filas", [])
        try:
            with transaction.atomic():
                # Revalidación dentro de transacción para evitar dobles cargas.
                for fila in filas:
                    item = Model.objects.filter(pk=fila["item_id"], activo=True).first()
                    if not item or (item.codigo or "").strip() != fila["codigo"]:
                        raise ValueError(f"El P/N {fila['codigo']} cambió o ya no está activo.")
                    stock = _get_or_create_stock(empresa, item)
                    stock = InventoryStock.objects.select_for_update().get(pk=stock.pk)
                    if stock.movimientos.exists() or stock.cantidad_fisica != 0:
                        raise ValueError(f"El P/N {fila['codigo']} ya tiene movimientos de inventario.")
                    cantidad = Decimal(fila["cantidad"])
                    costo = Decimal(fila["costo"])
                    stock.cantidad_fisica = cantidad
                    stock.costo_promedio = costo
                    stock.save(update_fields=["cantidad_fisica", "costo_promedio", "actualizado_en"])
                    InventoryMovement.objects.create(
                        stock=stock,
                        tipo=InventoryMovement.Tipo.INVENTARIO_INICIAL,
                        cantidad=cantidad,
                        costo_unitario=costo,
                        saldo_anterior=Decimal("0"),
                        saldo_nuevo=cantidad,
                        costo_promedio_anterior=Decimal("0"),
                        costo_promedio_nuevo=costo,
                        referencia="INVENTARIO INICIAL MASIVO",
                        motivo=fila.get("observacion") or "Carga masiva de inventario inicial",
                        creado_por=request.user,
                    )
        except ValueError as exc:
            messages.error(request, str(exc) + " No se cargó ningún artículo.")
            return redirect(f"{request.path}?empresa={empresa}")
        request.session.pop(session_key, None)
        messages.success(request, f"Inventario inicial cargado correctamente: {len(filas)} artículos.")
        return redirect(f"/inventario/existencias/?empresa={empresa}")

    if request.method == "POST":
        archivo = request.FILES.get("archivo")
        if not archivo:
            messages.error(request, "Seleccione un archivo Excel (.xlsx).")
            return redirect(f"{request.path}?empresa={empresa}")
        if not archivo.name.lower().endswith(".xlsx"):
            messages.error(request, "El archivo debe ser formato .xlsx.")
            return redirect(f"{request.path}?empresa={empresa}")
        try:
            from openpyxl import load_workbook
            wb = load_workbook(archivo, data_only=True, read_only=True)
            ws = wb["Inventario inicial"] if "Inventario inicial" in wb.sheetnames else wb.active
        except Exception as exc:
            messages.error(request, f"No fue posible leer el Excel: {exc}")
            return redirect(f"{request.path}?empresa={empresa}")

        items_por_codigo = {
            (i.codigo or "").strip(): i
            for i in Model.objects.filter(activo=True)
            if (i.codigo or "").strip()
        }
        filas_validas, errores, advertencias = [], [], []
        vistos = set()
        total_valor = Decimal("0")

        for nro, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            codigo = str(row[0] or "").strip()
            cantidad_raw = row[3] if len(row) > 3 else None
            costo_raw = row[4] if len(row) > 4 else None
            observacion = str(row[5] or "").strip() if len(row) > 5 else ""
            # Filas sin cantidad son deliberadamente ignoradas.
            if cantidad_raw is None or str(cantidad_raw).strip() == "":
                continue
            if not codigo:
                errores.append(f"Fila {nro}: falta P/N.")
                continue
            if codigo in vistos:
                errores.append(f"Fila {nro}: P/N duplicado en el archivo: {codigo}.")
                continue
            vistos.add(codigo)
            item = items_por_codigo.get(codigo)
            if not item:
                errores.append(f"Fila {nro}: P/N {codigo} no existe o está inactivo en el catálogo de destino.")
                continue
            try:
                cantidad = Decimal(str(cantidad_raw).replace(",", "."))
                costo = Decimal(str(costo_raw or 0).replace(",", "."))
            except Exception:
                errores.append(f"Fila {nro}: cantidad o costo inválido para {codigo}.")
                continue
            if cantidad < 0 or costo < 0:
                errores.append(f"Fila {nro}: cantidad y costo no pueden ser negativos ({codigo}).")
                continue
            # Cero explícito se ignora: no crea un movimiento inicial innecesario.
            if cantidad == 0:
                advertencias.append(f"Fila {nro}: {codigo} tiene cantidad 0 y será omitido.")
                continue
            stock = InventoryStock.objects.filter(
                empresa=empresa, catalogo=catalogo, catalogo_item_id=item.pk
            ).first()
            if stock and (stock.movimientos.exists() or stock.cantidad_fisica != 0):
                errores.append(f"Fila {nro}: {codigo} ya tiene inventario/movimientos. Use Ajuste de inventario.")
                continue
            if costo == 0:
                advertencias.append(f"Fila {nro}: {codigo} se cargará con costo $0.")
            valor = cantidad * costo
            total_valor += valor
            filas_validas.append({
                "fila": nro,
                "item_id": item.pk,
                "codigo": codigo,
                "descripcion": item.descripcion or "",
                "unidad": item.unidad_medida or "UND",
                "cantidad": str(cantidad),
                "costo": str(costo),
                "valor": str(valor),
                "observacion": observacion,
            })

        datos_session = {
            "filas": filas_validas,
            "errores": errores,
            "advertencias": advertencias,
            "total_valor": str(total_valor),
        }
        request.session[session_key] = datos_session
        request.session.modified = True
        contexto["preview"] = {
            "filas": filas_validas[:100],
            "total_filas": len(filas_validas),
            "errores": errores,
            "advertencias": advertencias,
            "total_valor": total_valor,
            "truncado": len(filas_validas) > 100,
        }

    return render(request, "inventario/inventario_inicial_masivo.html", contexto)

@login_required
@inventario_required
def reserva_transicion(request):
    """Reserva manual de apertura para PAW que ya estaban activos al iniciar Etapa 2."""
    from compras_oil.models import PurchaseRequest
    from .models import InventoryStock, InventoryReservation

    stocks = InventoryStock.objects.filter(
        empresa=InventoryStock.Empresa.IMPETUS,
        cantidad_fisica__gt=F("cantidad_reservada"),
    ).order_by("codigo")

    if request.method == "POST":
        paw_numero = (request.POST.get("paw_numero") or "").strip()
        stock_id = request.POST.get("stock_id")
        raw = (request.POST.get("cantidad") or "0").strip()
        observacion = (request.POST.get("observacion") or "").strip()
        try:
            cantidad = Decimal(raw.replace(",", "."))
        except Exception:
            cantidad = Decimal("0")

        compra = PurchaseRequest.objects.filter(paw_numero__iexact=paw_numero).order_by("-id").first()
        if not compra:
            messages.error(request, f"No se encontró un PAW #{paw_numero} con solicitud/BOM en el sistema.")
        elif cantidad <= 0:
            messages.error(request, "La cantidad a reservar debe ser mayor que cero.")
        else:
            with transaction.atomic():
                stock = get_object_or_404(InventoryStock.objects.select_for_update(), pk=stock_id)
                if stock.empresa != InventoryStock.Empresa.IMPETUS:
                    messages.error(request, "La reserva de transición debe salir del inventario IMPETUS.")
                elif cantidad > Decimal(stock.cantidad_disponible):
                    messages.error(request, f"Disponible real: {stock.cantidad_disponible}.")
                else:
                    InventoryReservation.objects.create(
                        stock=stock, cantidad=cantidad, purchase_request=compra,
                        es_transicion=True, creado_por=request.user,
                        observacion=observacion or f"Reserva inicial de transición PAW #{compra.paw_numero}",
                    )
                    stock.cantidad_reservada = F("cantidad_reservada") + cantidad
                    stock.save(update_fields=["cantidad_reservada", "actualizado_en"])
                    messages.success(request, f"Reservadas {cantidad} unidades de {stock.codigo} para PAW #{compra.paw_numero}. La existencia física no cambió.")
                    return redirect("inventario:reserva_transicion")

    reservas = InventoryReservation.objects.filter(es_transicion=True).select_related("stock", "purchase_request", "creado_por").order_by("-creado_en")[:100]
    return render(request, "inventario/reserva_transicion.html", {"stocks": stocks, "reservas": reservas})
